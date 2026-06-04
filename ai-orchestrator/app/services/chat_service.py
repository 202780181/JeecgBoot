from collections.abc import AsyncIterator
import base64
from io import BytesIO
import json
from json import JSONDecodeError
from pathlib import Path
import re
from urllib.parse import urlparse
from uuid import uuid4

import httpx

from app.core.config import settings
from app.models.event_display import operation_display
from app.models.events import StreamEventType, StreamEventWriter
from app.models.schemas import AppChatStreamRequest, ModelConfig
from app.services.agent_loop import BuilderAgentLoop
from app.services.agent_preflight import AgentIntentClassification, AgentPreflight, AgentPreflightResult, PreflightIntent
from app.services.context_builder import ContextBuilder
from app.services.jeecg_client import JeecgClient
from app.services.spec_service import SpecKitResult, generate_spec_with_speckit
from app.skills import SkillRegistry
from app.skills.base import SkillSpec
from app.tools import ToolRegistry
from app.tools.base import BaseTool, ToolResult


class SkillValidationError(ValueError):
    def __init__(
        self,
        message: str,
        code: str,
        *,
        skill_ids: list[str] | None = None,
        tool_names: list[str] | None = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.skill_ids = skill_ids or []
        self.tool_names = tool_names or []

    def to_event_data(self) -> dict:
        data = {
            "code": self.code,
            "message": str(self),
        }
        if self.skill_ids:
            data["skillIds"] = self.skill_ids
        if self.tool_names:
            data["toolNames"] = self.tool_names
        return data


class ChatService:
    def __init__(
        self,
        jeecg_client: JeecgClient | None = None,
        tool_registry: ToolRegistry | None = None,
        skill_registry: SkillRegistry | None = None,
        context_builder: ContextBuilder | None = None,
        agent_preflight: AgentPreflight | None = None,
        spec_generator=generate_spec_with_speckit,
    ) -> None:
        self.jeecg_client = jeecg_client or JeecgClient()
        self.tool_registry = tool_registry or ToolRegistry()
        self.skill_registry = skill_registry or SkillRegistry()
        self.context_builder = context_builder or ContextBuilder()
        self.agent_preflight = agent_preflight or AgentPreflight()
        self.spec_generator = spec_generator

    async def stream(self, request: AppChatStreamRequest) -> AsyncIterator[str]:
        run_id = request.run_id
        conversation_id = request.conversation_id or "debug"
        topic_id = request.topic_id or ""
        event_writer = StreamEventWriter(run_id=run_id, conversation_id=conversation_id, topic_id=topic_id)
        yield event_writer.event(
            StreamEventType.RUN_STARTED,
            {"runId": run_id},
        )
        try:
            selected_skills = self._resolve_and_validate_skills(request.skill_ids)
            model = await self.jeecg_client.get_model_config(
                request.app.model_id,
                request.user_context,
            )
            available_tool_names = self._available_tool_names(request.skill_ids, request.enable_search, request)
            classification = await self._classify_agent_intent(model, request)
            preflight = self.agent_preflight.analyze(request, available_tool_names, classification)
            setattr(request, "_allowed_tool_names", preflight.allowed_tools)
            yield event_writer.event(
                StreamEventType.PREFLIGHT,
                {
                    **preflight.to_event_data(),
                    "title": "正在理解需求",
                    "summary": preflight.operation_note or preflight.summary,
                    "display": operation_display(
                        preflight.operation_note or preflight.summary or "正在理解需求并检查操作安全",
                        detail="；".join(preflight.proposed_steps[:3]),
                        icon="search",
                        status="error" if preflight.blocked or preflight.requires_confirmation else "done",
                    ),
                },
            )
            if preflight.blocked or preflight.needs_clarification or preflight.requires_confirmation:
                yield event_writer.event(
                    StreamEventType.MESSAGE,
                    {"message": self._preflight_stop_message(preflight)},
                )
                yield event_writer.event(
                    StreamEventType.MESSAGE_END,
                    None,
                )
                return
            for skill in selected_skills:
                yield event_writer.event(
                    StreamEventType.SKILL_SELECTED,
                    {
                        "skillId": skill.id,
                        "skillName": skill.name,
                        "description": skill.description,
                        "category": skill.category,
                        "selectionMode": skill.selection_mode,
                        "availableToolNames": skill.available_tool_names,
                        "defaultToolNames": skill.default_tool_names,
                        "forbiddenToolNames": skill.forbidden_tool_names,
                        "title": f"正在执行 Skill：{skill.name}",
                        "summary": skill.description,
                        "display": operation_display(f"正在执行 Skill：{skill.name}", detail=skill.description, icon="search", status="running"),
                    },
                )
                if skill.spec_kit.enabled:
                    async for event in self._run_skill_spec_kit(
                        event_writer,
                        skill,
                        request.input,
                    ):
                        yield event
            messages = await self._initial_messages(request, model)
            context_metadata = request.context_source.summary.metadata or {}
            context_selection = context_metadata.get("contextSelection")
            attachment_selection = context_metadata.get("attachmentSelection")
            if self._should_emit_context_selection(context_selection):
                yield event_writer.event(
                    StreamEventType.CONTEXT_SELECTED,
                    {
                        "contextSelection": context_selection,
                        "attachmentSelection": attachment_selection,
                        "title": "已选择上下文",
                        "summary": "已完成上下文预算与片段选择",
                        "display": operation_display("已选择上下文", detail="已完成上下文预算与片段选择", icon="search", status="done"),
                    },
                )
            agent_loop = BuilderAgentLoop(
                tool_resolver=self._resolve_tool_call,
                tool_call_id=self._tool_call_id,
                assistant_tool_call_message=self._assistant_tool_call_message,
                tool_result_message=self._tool_result_message,
                stream_complete=self._stream_complete,
                max_repair_attempts=self._max_repair_attempts(model),
                allowed_tool_names=preflight.allowed_tools,
            )
            async for event in agent_loop.run(
                model=model,
                request=request,
                messages=messages,
                event_writer=event_writer,
                max_tool_rounds=self._max_tool_rounds(model, selected_skills),
            ):
                yield event

            yield event_writer.event(
                StreamEventType.MESSAGE_END,
                None,
            )
        except SkillValidationError as exc:
            error_data = exc.to_event_data()
            error_data["title"] = "执行失败"
            error_data["summary"] = str(exc)
            error_data["display"] = operation_display("执行失败", detail=str(exc), icon="search", status="error")
            yield event_writer.event(
                StreamEventType.ERROR,
                error_data,
            )
        except Exception as exc:
            message = f"AI 对话失败：{exc}"
            yield event_writer.event(
                StreamEventType.ERROR,
                {
                    "message": message,
                    "title": "执行失败",
                    "summary": message,
                    "display": operation_display("执行失败", detail=message, icon="search", status="error"),
                },
            )

    async def _stream_complete(
        self,
        model: ModelConfig,
        request: AppChatStreamRequest,
        messages: list[dict],
        use_tools: bool = False,
    ) -> AsyncIterator[str | dict]:
        payload, endpoint, headers, timeout = self._build_model_request(
            model,
            request,
            stream=True,
            messages=messages,
            use_tools=use_tools,
        )
        saw_stream_chunk = False
        pending_tool_calls: dict[int, dict] = {}
        async with httpx.AsyncClient(timeout=timeout) as client:
            async with client.stream("POST", endpoint, json=payload, headers=headers) as response:
                if not response.is_success:
                    body = await response.aread()
                    self._raise_for_bad_model_response(
                        httpx.Response(
                            response.status_code,
                            content=body,
                            headers=response.headers,
                            request=response.request,
                        ),
                        endpoint,
                    )
                content_type = response.headers.get("content-type", "")
                if "text/event-stream" not in content_type.lower():
                    body = await response.aread()
                    completion = self._parse_model_response(
                        httpx.Response(
                            response.status_code,
                            content=body,
                            headers=response.headers,
                            request=response.request,
                        ),
                        endpoint,
                    )
                    tool_calls = self._extract_tool_calls(completion)
                    if tool_calls:
                        yield tool_calls[0]
                        return
                    yield self._extract_completion_text(completion, endpoint)
                    return

                async for line in response.aiter_lines():
                    if not line.startswith("data:"):
                        continue
                    data = line.removeprefix("data:").strip()
                    if not data:
                        continue
                    if data == "[DONE]":
                        break
                    chunk, tool_calls = self._parse_model_stream_chunk(data, endpoint)
                    if chunk:
                        saw_stream_chunk = True
                        yield chunk
                    for tool_call_delta in tool_calls:
                        self._merge_tool_call_delta(pending_tool_calls, tool_call_delta)

                tool_calls = self._finalize_tool_calls(pending_tool_calls)
                if tool_calls:
                    yield tool_calls[0]
                    return
                if not saw_stream_chunk:
                    raise ValueError("模型未返回流式文本")

    async def _complete(self, model: ModelConfig, request: AppChatStreamRequest) -> str:
        payload, endpoint, headers, timeout = self._build_model_request(
            model,
            request,
            stream=False,
            messages=await self._initial_messages(request, model),
            use_tools=False,
        )
        async with httpx.AsyncClient(timeout=timeout) as client:
            response = await client.post(endpoint, json=payload, headers=headers)
        self._raise_for_bad_model_response(response, endpoint)
        return self._extract_completion_text(self._parse_model_response(response, endpoint), endpoint)

    def _should_emit_context_selection(self, context_selection: dict | None) -> bool:
        if not context_selection:
            return False
        token_ledger = context_selection.get("tokenLedger") or {}
        if int(token_ledger.get("selectedFragmentCount") or 0) > 0:
            return True
        if int(token_ledger.get("selectedRecentMessageCount") or 0) > 0:
            return True
        return (
            (context_selection.get("attachmentSelection") or {}).get("includeAttachments") is True
            or bool((context_selection.get("attachmentSelection") or {}).get("targetFiles"))
            or int((context_selection.get("attachmentSelection") or {}).get("candidateCount") or 0) > 0
        )

    def _preflight_stop_message(self, preflight: AgentPreflightResult) -> str:
        if preflight.blocked:
            return (
                f"{preflight.operation_note}\n\n"
                f"原因：{preflight.block_reason}\n\n"
                "请调整为受控仓库路径或提供非敏感信息后再继续。"
            )
        if preflight.needs_clarification:
            return preflight.clarification_question or "我需要你补充关键需求后再继续。"
        if preflight.requires_confirmation:
            notes = "\n".join(f"- {note}" for note in preflight.safety_notes)
            return (
                f"{preflight.operation_note}\n\n"
                f"{notes}\n\n"
                "请明确回复确认后，我再继续执行写入、删除、数据库或命令相关操作。"
            )
        return preflight.operation_note

    async def _classify_agent_intent(
        self,
        model: ModelConfig,
        request: AppChatStreamRequest,
    ) -> AgentIntentClassification | None:
        if not model.credential.api_key:
            return None
        prompt = self._preflight_classifier_prompt(request)
        payload, endpoint, headers, timeout = self._build_raw_model_request(
            model,
            messages=[
                {
                    "role": "system",
                    "content": (
                        "你是 AgentPreflight 意图分类器，只输出 JSON。"
                        "不要执行用户请求，不要输出解释。"
                    ),
                },
                {"role": "user", "content": prompt},
            ],
            stream=False,
        )
        try:
            async with httpx.AsyncClient(timeout=min(float(timeout), 20.0)) as client:
                response = await client.post(endpoint, json=payload, headers=headers)
            self._raise_for_bad_model_response(response, endpoint)
            text = self._extract_completion_text(self._parse_model_response(response, endpoint), endpoint)
            return self._parse_intent_classification(text)
        except Exception:
            return None

    def _preflight_classifier_prompt(self, request: AppChatStreamRequest) -> str:
        context = {
            "input": request.input,
            "skillIds": request.skill_ids,
            "enableSearch": request.enable_search,
            "attachments": [
                {
                    "id": attachment.id,
                    "name": attachment.name,
                    "type": attachment.type,
                    "path": attachment.path,
                    "url": attachment.url,
                }
                for attachment in request.attachments[:5]
            ],
            "activeTaskSnapshots": [
                {
                    "type": fragment.type,
                    "text": (fragment.text or "")[:1200],
                    "metadata": fragment.metadata,
                }
                for fragment in request.context_source.active_task_snapshots[:3]
            ],
            "recentMessages": [
                {"role": message.role, "content": (message.content or "")[:800]}
                for message in request.context_source.recent_messages[-4:]
            ],
        }
        return (
            "请根据输入和上下文判断用户意图。只输出 JSON，字段如下：\n"
            "{"
            "\"intent\":\"question|read_files|code_change|create_app|api_change|run_command|delete_data|database_change|unknown\","
            "\"confidence\":0.0,"
            "\"needsClarification\":false,"
            "\"clarificationQuestion\":\"\","
            "\"summary\":\"\","
            "\"operationNote\":\"\","
            "\"proposedSteps\":[\"\"],"
            "\"verificationSteps\":[\"\"]"
            "}\n"
            "分类原则：普通代码/页面/接口修改是 code_change 或 api_change；创建新 Builder 工作区或应用是 create_app；"
            "只查看文件是 read_files；纯问答是 question；数据库删除/清空是 delete_data 或 database_change。\n"
            f"输入上下文：{json.dumps(context, ensure_ascii=False)}"
        )

    def _parse_intent_classification(self, text: str) -> AgentIntentClassification | None:
        body = self._extract_json_object(text)
        if not body:
            return None
        intent = self._parse_preflight_intent(body.get("intent"))
        return AgentIntentClassification(
            intent=intent,
            confidence=self._float_value(body.get("confidence")),
            needs_clarification=bool(body.get("needsClarification") or body.get("needs_clarification")),
            clarification_question=str(body.get("clarificationQuestion") or body.get("clarification_question") or ""),
            summary=str(body.get("summary") or ""),
            operation_note=str(body.get("operationNote") or body.get("operation_note") or ""),
            proposed_steps=self._string_list(body.get("proposedSteps") or body.get("proposed_steps")),
            verification_steps=self._string_list(body.get("verificationSteps") or body.get("verification_steps")),
        )

    def _extract_json_object(self, text: str) -> dict | None:
        value = (text or "").strip()
        if value.startswith("```"):
            value = re.sub(r"^```(?:json)?\s*", "", value)
            value = re.sub(r"\s*```$", "", value)
        start = value.find("{")
        end = value.rfind("}")
        if start < 0 or end < start:
            return None
        try:
            body = json.loads(value[start : end + 1])
        except JSONDecodeError:
            return None
        return body if isinstance(body, dict) else None

    def _parse_preflight_intent(self, value: object) -> PreflightIntent:
        try:
            return PreflightIntent(str(value or "unknown"))
        except ValueError:
            return PreflightIntent.UNKNOWN

    def _float_value(self, value: object) -> float:
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0

    def _string_list(self, value: object) -> list[str]:
        if not isinstance(value, list):
            return []
        return [str(item) for item in value if str(item).strip()][:8]

    async def _run_skill_spec_kit(
        self,
        event_writer: StreamEventWriter,
        skill: SkillSpec,
        requirement: str,
    ) -> AsyncIterator[str]:
        yield self._spec_event(
            event_writer,
            skill,
            stage="start",
            status="running",
            message=f"开始执行 spec-kit：{skill.spec_kit.template or skill.name}",
        )
        try:
            result = self.spec_generator(requirement)
        except Exception as exc:
            yield self._spec_event(
                event_writer,
                skill,
                stage="failed",
                status="failed",
                message=f"spec-kit 执行失败：{exc}",
            )
            return

        for stage, title, path in (
            ("spec", "Spec 已生成", result.spec_path),
            ("plan", "Plan 已生成", result.plan_path),
            ("tasks", "Tasks 已生成", result.tasks_path),
        ):
            yield self._spec_event(
                event_writer,
                skill,
                stage=stage,
                status="completed",
                message=title,
                artifact={
                    "type": stage,
                    "path": path,
                    "title": title,
                },
            )
        yield self._spec_event(
            event_writer,
            skill,
            stage="completed",
            status="completed",
            message="spec-kit 已完成 Spec / Plan / Tasks 生成",
            result=result,
        )

    def _spec_event(
        self,
        event_writer: StreamEventWriter,
        skill: SkillSpec,
        *,
        stage: str,
        status: str,
        message: str,
        artifact: dict | None = None,
        result: SpecKitResult | None = None,
    ) -> str:
        data = {
            "stage": stage,
            "status": status,
            "message": message,
            "skillId": skill.id,
            "skillName": skill.name,
            "template": skill.spec_kit.template,
            "mode": skill.spec_kit.mode,
            "title": message,
            "summary": skill.name,
            "display": operation_display(message, detail=skill.name, icon="search", status="running" if status == "running" else "done" if status == "completed" else "error"),
        }
        if artifact:
            data["artifact"] = artifact
        if result:
            data["result"] = {
                "note": result.note,
                "source": result.source,
                "featureDir": result.feature_dir,
                "specPath": result.spec_path,
                "planPath": result.plan_path,
                "tasksPath": result.tasks_path,
                "taskCount": len(result.tasks),
            }
        return event_writer.event(
            StreamEventType.SPEC_EVENT,
            data,
        )

    async def _initial_messages(self, request: AppChatStreamRequest, model: ModelConfig | None = None) -> list[dict]:
        prompt = request.app.prompt or "你是 JeecgBoot AI 应用开发助手。请直接、准确地回答用户问题。"
        skill_prompt = self._skill_prompt(request.skill_ids)
        if skill_prompt:
            prompt = f"{prompt}\n\n{skill_prompt}"
        if request.enable_search:
            prompt = (
                f"{prompt}\n\n"
                "联网搜索已开启。遇到实时信息、新闻、价格、版本、官方文档、网页资料或你不确定的信息时，"
                "优先调用 web_search 工具；回答时要基于搜索结果总结，并尽量附上来源链接。"
            )
        current_user_content = await self._current_user_content_with_attachments(request)
        if model is None:
            model = ModelConfig(id="default", model_name="default", base_url="http://localhost")
        return await self.context_builder.build(request, model, prompt, current_user_content)

    async def _current_user_content_with_attachments(self, request: AppChatStreamRequest):
        text = await self._user_input_with_attachments(request)
        image_parts = await self._image_content_parts(request.attachments)
        if not image_parts:
            return text
        return [{"type": "text", "text": text}, *image_parts]

    async def _user_input_with_attachments(self, request: AppChatStreamRequest) -> str:
        if not request.attachments:
            return request.input
        extracted_attachments = await self._extract_attachments(request)
        lines = [request.input, "", "用户上传了以下附件，请优先结合可读取的附件正文回答："]
        for index, attachment in enumerate(request.attachments, start=1):
            name = attachment.name or "未命名附件"
            file_type = attachment.type or "unknown"
            size = attachment.size if attachment.size is not None else "unknown"
            url = attachment.url or attachment.path or ""
            lines.append(f"{index}. {name} | type={file_type} | size={size} | url={url}")
            if self._is_image_attachment(attachment):
                lines.append("图片附件已作为视觉输入传给模型。")
                continue
            attachment_text = extracted_attachments.get(index - 1, "")
            if attachment_text:
                lines.append("```text")
                lines.append(attachment_text)
                lines.append("```")
            else:
                status = attachment.extraction_status or "unreadable"
                error = attachment.extraction_error or "当前文件类型暂未解析出正文"
                lines.append(f"附件读取状态：{status}，{error}")
        return "\n".join(lines)

    async def _extract_attachments(self, request: AppChatStreamRequest) -> dict[int, str]:
        extracted: dict[int, str] = {}
        for index, attachment in enumerate(request.attachments):
            if self._is_image_attachment(attachment):
                attachment.extraction_status = "vision_input"
                attachment.extraction_error = None
                continue
            if attachment.extracted_text:
                text = self._limit_attachment_text(attachment.extracted_text)
                attachment.extracted_text = text
                attachment.extraction_status = "provided"
                extracted[index] = text
                continue
            try:
                text = await self._extract_attachment_text(attachment)
            except Exception as exc:
                attachment.extraction_status = "failed"
                attachment.extraction_error = str(exc)
                continue
            if text:
                text = self._limit_attachment_text(text)
                attachment.extracted_text = text
                attachment.extraction_status = "extracted"
                extracted[index] = text
            else:
                attachment.extraction_status = "empty"
                attachment.extraction_error = "未读取到可用文本"
        return extracted

    async def _extract_attachment_text(self, attachment) -> str:
        url = attachment.url or ""
        if not url:
            raise ValueError("附件缺少可读取 URL")
        content = await self._download_attachment(url)
        suffix = self._attachment_suffix(attachment)
        if suffix in {".txt", ".md", ".json", ".csv", ".xml", ".yaml", ".yml", ".log"}:
            return self._decode_text(content)
        if suffix == ".pdf":
            return self._extract_pdf_text(content)
        if suffix == ".docx":
            return self._extract_docx_text(content)
        if suffix == ".xlsx":
            return self._extract_xlsx_text(content)
        raise ValueError(f"暂不支持解析 {suffix or attachment.type or 'unknown'} 文件正文")

    async def _image_content_parts(self, attachments) -> list[dict]:
        parts: list[dict] = []
        for attachment in attachments:
            if not self._is_image_attachment(attachment):
                continue
            image_url = await self._image_data_url(attachment)
            if not image_url:
                image_url = attachment.url or attachment.path
            if not image_url:
                continue
            parts.append(
                {
                    "type": "image_url",
                    "image_url": {
                        "url": image_url,
                        "detail": "auto",
                    },
                }
            )
        return parts

    async def _image_data_url(self, attachment) -> str | None:
        url = attachment.url or ""
        if url.startswith("data:image/"):
            return url
        if not url:
            return None
        try:
            content = await self._download_attachment(url)
        except Exception:
            return None
        if not content:
            return None
        if len(content) > 20 * 1024 * 1024:
            return None
        mime_type = attachment.type if (attachment.type or "").startswith("image/") else self._image_mime_type(attachment)
        encoded = base64.b64encode(content).decode("ascii")
        return f"data:{mime_type};base64,{encoded}"

    def _is_image_attachment(self, attachment) -> bool:
        file_type = (attachment.type or "").lower()
        if file_type.startswith("image/"):
            return True
        return self._attachment_suffix(attachment) in {".png", ".jpg", ".jpeg", ".webp", ".gif"}

    def _image_mime_type(self, attachment) -> str:
        suffix = self._attachment_suffix(attachment)
        if suffix in {".jpg", ".jpeg"}:
            return "image/jpeg"
        if suffix == ".webp":
            return "image/webp"
        if suffix == ".gif":
            return "image/gif"
        return "image/png"

    async def _download_attachment(self, url: str) -> bytes:
        async with httpx.AsyncClient(timeout=20, follow_redirects=True) as client:
            response = await client.get(url)
        response.raise_for_status()
        return response.content

    def _attachment_suffix(self, attachment) -> str:
        name = attachment.name or attachment.path or attachment.url or ""
        return Path(urlparse(name).path).suffix.lower()

    def _decode_text(self, content: bytes) -> str:
        for encoding in ("utf-8", "utf-8-sig", "gb18030"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="ignore")

    def _extract_pdf_text(self, content: bytes) -> str:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        texts = []
        for page in reader.pages[:20]:
            texts.append(page.extract_text() or "")
        return "\n".join(texts)

    def _extract_docx_text(self, content: bytes) -> str:
        from docx import Document

        document = Document(BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)

    def _extract_xlsx_text(self, content: bytes) -> str:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets[:5]:
            lines.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(max_row=200, values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    lines.append("\t".join(values))
        workbook.close()
        return "\n".join(lines)

    def _limit_attachment_text(self, text: str, limit: int = 20000) -> str:
        normalized = text.replace("\x00", "").strip()
        if len(normalized) <= limit:
            return normalized
        return f"{normalized[:limit]}\n\n[附件正文已截断，仅保留前 {limit} 字符]"

    def _context_messages(self, request: AppChatStreamRequest) -> list[dict]:
        context_messages: list[dict] = []
        for message in request.messages[-12:]:
            role = message.role if message.role in {"user", "assistant"} else ""
            content = message.content.strip()
            if not role or not content or content == request.input:
                continue
            context_messages.append({"role": role, "content": content[:4000]})
        return context_messages

    def _skill_prompt(self, skill_ids: list[str] | None) -> str:
        skills = self.skill_registry.resolve(skill_ids)
        if not skills:
            return ""
        instructions = "\n".join(f"- {skill.name}: {skill.instruction}" for skill in skills)
        spec_instructions = "\n".join(
            f"- {skill.name}: template={skill.spec_kit.template or 'default'}，mode={skill.spec_kit.mode}。{skill.spec_kit.instruction or ''}"
            for skill in skills
            if skill.spec_kit.enabled
        )
        prompt = f"已启用以下 Skills，请严格结合这些技能约束回答：\n{instructions}"
        if spec_instructions:
            prompt = f"{prompt}\n\nspec-kit 绑定策略：\n{spec_instructions}"
        return prompt

    def _resolve_and_validate_skills(self, skill_ids: list[str] | None) -> list[SkillSpec]:
        normalized_skill_ids = self._dedupe_names(skill_ids or [])
        if not normalized_skill_ids:
            return []

        selected_skills: list[SkillSpec] = []
        for skill_id in normalized_skill_ids:
            try:
                selected_skills.append(self.skill_registry.get(skill_id))
            except ValueError as exc:
                raise SkillValidationError(
                    f"Skill 不存在或未注册：{skill_id}",
                    "SKILL_NOT_FOUND",
                    skill_ids=[skill_id],
                ) from exc

        self._validate_skill_conflicts(selected_skills)
        self._validate_skill_tools(selected_skills)
        return selected_skills

    def _validate_skill_conflicts(self, skills: list[SkillSpec]) -> None:
        if len(skills) <= 1:
            return
        single_skills = [skill for skill in skills if skill.selection_mode == "single"]
        if not single_skills:
            return
        raise SkillValidationError(
            "已选择的 Skills 存在互斥关系，请只保留一个单选 Skill 后再发送。",
            "SKILL_CONFLICT",
            skill_ids=[skill.id for skill in skills],
        )

    def _validate_skill_tools(self, skills: list[SkillSpec]) -> None:
        registered_tool_names = {tool.spec.name for tool in self.tool_registry.list_tools()}
        missing_tool_names: list[str] = []
        for skill in skills:
            referenced_tool_names = [
                *(skill.available_tool_names or []),
                *skill.default_tool_names,
                *skill.forbidden_tool_names,
            ]
            for tool_name in referenced_tool_names:
                if tool_name not in registered_tool_names and tool_name not in missing_tool_names:
                    missing_tool_names.append(tool_name)
        if missing_tool_names:
            raise SkillValidationError(
                f"Skill 依赖的工具不可用：{', '.join(missing_tool_names)}",
                "SKILL_TOOL_UNAVAILABLE",
                skill_ids=[skill.id for skill in skills],
                tool_names=missing_tool_names,
            )

    def _available_tool_names(
        self,
        skill_ids: list[str] | None,
        enable_search: bool = False,
        request: AppChatStreamRequest | None = None,
    ) -> list[str] | None:
        skills = self.skill_registry.resolve(skill_ids)
        if not skills:
            if request and (self._should_enable_builder_tools(request) or self._looks_like_builder_request(request)):
                tool_names = self._builder_tool_names()
                if enable_search:
                    tool_names.append("web_search")
                return tool_names
            if enable_search:
                tool_names = self._non_search_tool_names()
                if "web_search" not in tool_names:
                    tool_names.append("web_search")
                return tool_names
            return self._non_search_tool_names()
        tool_names: list[str] | None = None
        forbidden_names: set[str] = set()
        has_explicit_tool_allowlist = False
        for skill in skills:
            forbidden_names.update(skill.forbidden_tool_names)
            if skill.available_tool_names is not None:
                has_explicit_tool_allowlist = True
                tool_names = tool_names or []
                for name in skill.available_tool_names:
                    if name not in tool_names:
                        tool_names.append(name)
            for name in skill.default_tool_names:
                tool_names = tool_names or []
                if name not in tool_names:
                    tool_names.append(name)
        if tool_names is None:
            tool_names = [tool.spec.name for tool in self.tool_registry.list_tools()]
        if enable_search and not has_explicit_tool_allowlist and "web_search" not in tool_names:
            tool_names.append("web_search")
        if not enable_search:
            forbidden_names.add("web_search")
        return [name for name in tool_names if name not in forbidden_names]

    def _non_search_tool_names(self) -> list[str]:
        return [
            tool.spec.name
            for tool in self.tool_registry.list_tools()
            if tool.spec.name != "web_search" and not tool.spec.name.startswith("builder_")
        ]

    def _builder_tool_names(self) -> list[str]:
        return [
            tool.spec.name
            for tool in self.tool_registry.list_tools()
            if tool.spec.name.startswith("builder_")
        ]

    def _should_enable_builder_tools(self, request: AppChatStreamRequest) -> bool:
        for text in self._builder_context_texts(request):
            if "ai-builder-workspaces" in text:
                return True
            for workspace_id in self._workspace_id_candidates(text):
                if self._builder_workspace_exists(workspace_id):
                    return True
        for fragment in [
            *request.context_source.active_task_snapshots,
            *request.context_source.relevant_fragments,
            *request.context_source.attachment_matches,
            *request.context_source.attachment_candidates,
            *request.context_source.attachment_summaries,
        ]:
            metadata = fragment.metadata or {}
            workspace_id = metadata.get("workspaceId")
            if isinstance(workspace_id, str) and self._builder_workspace_exists(workspace_id):
                return True
            if fragment.type in {"workspace_snapshot", "file_change", "build_result", "preview_url"}:
                return True
            for workspace_id in self._workspace_id_candidates(fragment.text or ""):
                if self._builder_workspace_exists(workspace_id):
                    return True
        return False

    def _looks_like_builder_request(self, request: AppChatStreamRequest) -> bool:
        text = (request.input or "").lower()
        keywords = (
            "修改",
            "新增",
            "写入",
            "重构",
            "实现",
            "修复",
            "改代码",
            "补充",
            "接入",
            "创建项目",
            "生成项目",
            "开发一个项目",
            "小程序",
            "h5",
            "uniapp",
            "页面",
            "接口",
        )
        return any(keyword.lower() in text for keyword in keywords)

    def _builder_context_texts(self, request: AppChatStreamRequest) -> list[str]:
        texts = [request.input or "", request.context_source.summary.text or ""]
        summary_metadata = request.context_source.summary.metadata or {}
        texts.extend(self._string_values(summary_metadata, max_depth=4))
        for fragment in request.context_source.active_task_snapshots:
            texts.append(fragment.text or "")
            texts.extend(self._string_values(fragment.metadata or {}, max_depth=4))
        for message in [*request.context_source.recent_messages, *request.context_source.relevant_messages]:
            texts.append(message.content or "")
            texts.extend(self._string_values(message.metadata or {}, max_depth=3))
        for message in request.messages:
            texts.append(message.content or "")
        return [text for text in texts if text]

    def _string_values(self, value: object, max_depth: int, depth: int = 0) -> list[str]:
        if depth > max_depth or value is None:
            return []
        if isinstance(value, str):
            return [value]
        if isinstance(value, dict):
            texts: list[str] = []
            for item in value.values():
                texts.extend(self._string_values(item, max_depth=max_depth, depth=depth + 1))
            return texts
        if isinstance(value, list):
            texts: list[str] = []
            for item in value:
                texts.extend(self._string_values(item, max_depth=max_depth, depth=depth + 1))
            return texts
        return []

    def _workspace_id_candidates(self, text: str) -> list[str]:
        return [
            token
            for token in re.findall(r"[a-zA-Z0-9][a-zA-Z0-9_-]{7,}", text or "")
            if "-" in token
        ]

    def _builder_workspace_exists(self, workspace_id: str) -> bool:
        try:
            from app.builder import BuilderTools

            BuilderTools().workspace_manager.get_workspace_path(workspace_id)
            return True
        except Exception:
            return False

    def _max_tool_rounds(self, model: ModelConfig, skills: list[SkillSpec]) -> int:
        params = model.model_params or {}
        for key in ("maxToolRounds", "max_tool_rounds", "toolRoundsLimit"):
            value = params.get(key)
            if value is None:
                continue
            try:
                return max(1, min(int(value), 128))
            except (TypeError, ValueError):
                continue
        skill_limits = [skill.max_tool_rounds for skill in skills if skill.max_tool_rounds]
        if skill_limits:
            return max(1, min(max(skill_limits), 128))
        return max(1, min(settings.max_tool_rounds, 128))

    def _max_repair_attempts(self, model: ModelConfig) -> int:
        params = model.model_params or {}
        for key in ("maxRepairAttempts", "max_repair_attempts", "builderRepairAttempts"):
            value = params.get(key)
            if value is None:
                continue
            try:
                return max(1, min(int(value), 8))
            except (TypeError, ValueError):
                continue
        return 3

    def _dedupe_names(self, names: list[str]) -> list[str]:
        result: list[str] = []
        for name in names:
            if name and name not in result:
                result.append(name)
        return result

    def _build_model_request(
        self,
        model: ModelConfig,
        request: AppChatStreamRequest,
        stream: bool,
        messages: list[dict],
        use_tools: bool,
        ) -> tuple[dict, str, dict, int | float]:
        if not model.base_url:
            raise ValueError("模型 API 域名不能为空")
        if not model.model_name:
            raise ValueError("模型名称不能为空")
        if not model.credential.api_key:
            raise ValueError("模型 API Key 不能为空")

        endpoint = self._chat_completions_endpoint(model.base_url)

        payload = {
            "model": model.model_name,
            "messages": messages,
            "stream": stream,
        }
        if use_tools:
            allowed_tool_names = getattr(request, "_allowed_tool_names", None)
            tool_names = allowed_tool_names if allowed_tool_names is not None else self._available_tool_names(request.skill_ids, request.enable_search, request)
            tools = self.tool_registry.openai_tools(tool_names)
            if tools:
                payload["tools"] = tools
                payload["tool_choice"] = "auto"
        self._merge_generation_params(payload, model.model_params)
        payload["stream"] = stream
        headers = {
            "Authorization": f"Bearer {model.credential.api_key}",
            "Content-Type": "application/json",
            "Accept": "text/event-stream" if stream else "application/json",
        }
        return payload, endpoint, headers, model.model_params.get("timeout") or 120

    def _build_raw_model_request(
        self,
        model: ModelConfig,
        *,
        messages: list[dict],
        stream: bool,
    ) -> tuple[dict, str, dict, int | float]:
        if not model.base_url:
            raise ValueError("模型 API 域名不能为空")
        if not model.model_name:
            raise ValueError("模型名称不能为空")
        if not model.credential.api_key:
            raise ValueError("模型 API Key 不能为空")
        payload = {
            "model": model.model_name,
            "messages": messages,
            "stream": stream,
            "temperature": 0,
        }
        endpoint = self._chat_completions_endpoint(model.base_url)
        headers = {
            "Authorization": f"Bearer {model.credential.api_key}",
            "Content-Type": "application/json",
            "Accept": "text/event-stream" if stream else "application/json",
        }
        return payload, endpoint, headers, model.model_params.get("timeout") or 120

    def _extract_completion_text(self, body: dict, endpoint: str) -> str:
        choices = body.get("choices") or []
        if not choices:
            raise ValueError(f"模型未返回内容：endpoint={endpoint}")
        message = choices[0].get("message") or {}
        content = message.get("content") or choices[0].get("text") or ""
        return content.strip() or "模型未返回有效文本。"

    def _extract_tool_calls(self, body: dict) -> list[dict]:
        choices = body.get("choices") or []
        if not choices:
            return []
        message = choices[0].get("message") or {}
        tool_calls = message.get("tool_calls") or []
        return tool_calls if isinstance(tool_calls, list) else []

    def _parse_model_stream_chunk(self, data: str, endpoint: str) -> tuple[str, list[dict]]:
        try:
            body = json.loads(data)
        except JSONDecodeError:
            raise ValueError(f"模型流式响应格式不正确：endpoint={endpoint}，响应={data[:300]}") from None
        choices = body.get("choices") or []
        if not choices:
            return "", []
        choice = choices[0]
        delta = choice.get("delta") or {}
        content = delta.get("content")
        if content is None:
            content = choice.get("text")
        tool_calls = delta.get("tool_calls") or []
        if not isinstance(tool_calls, list):
            tool_calls = []
        return (str(content) if content is not None else ""), tool_calls

    def _merge_tool_call_delta(self, pending_tool_calls: dict[int, dict], delta: dict) -> None:
        index = int(delta.get("index") or 0)
        pending = pending_tool_calls.setdefault(
            index,
            {
                "id": "",
                "type": "function",
                "function": {"name": "", "arguments": ""},
            },
        )
        if delta.get("id"):
            pending["id"] = delta["id"]
        function_delta = delta.get("function") or {}
        if function_delta.get("name"):
            pending["function"]["name"] += function_delta["name"]
        if function_delta.get("arguments"):
            pending["function"]["arguments"] += function_delta["arguments"]

    def _finalize_tool_calls(self, pending_tool_calls: dict[int, dict]) -> list[dict]:
        return [pending_tool_calls[index] for index in sorted(pending_tool_calls)]

    def _resolve_tool_call(self, tool_call: dict) -> tuple[BaseTool, dict]:
        function = tool_call.get("function") or {}
        name = function.get("name")
        if not name:
            raise ValueError("模型返回的工具调用缺少工具名称")
        arguments = function.get("arguments") or "{}"
        if isinstance(arguments, str):
            try:
                tool_input = json.loads(arguments or "{}")
            except JSONDecodeError:
                raise ValueError(f"模型返回的工具参数不是 JSON：tool={name}，arguments={arguments[:300]}") from None
        elif isinstance(arguments, dict):
            tool_input = arguments
        else:
            raise ValueError(f"模型返回的工具参数格式不正确：tool={name}")
        return self.tool_registry.get(str(name)), tool_input

    def _tool_call_id(self, tool_call: dict) -> str:
        return str(tool_call.get("id") or f"call-{uuid4()}")

    def _assistant_tool_call_message(self, tool_call: dict) -> dict:
        return {
            "role": "assistant",
            "content": None,
            "tool_calls": [tool_call],
        }

    def _tool_result_message(self, tool_call_id: str, tool_result: ToolResult) -> dict:
        return {
            "role": "tool",
            "tool_call_id": tool_call_id,
            "name": tool_result.tool_name,
            "content": json.dumps(tool_result.model_dump(by_alias=True), ensure_ascii=False),
        }

    def _raise_for_bad_model_response(self, response: httpx.Response, endpoint: str) -> None:
        if response.is_success:
            return
        detail = self._response_excerpt(response)
        raise ValueError(f"模型接口调用失败：HTTP {response.status_code}，endpoint={endpoint}，响应={detail}")

    def _parse_model_response(self, response: httpx.Response, endpoint: str) -> dict:
        if not response.content:
            raise ValueError(f"模型接口返回空响应：endpoint={endpoint}")
        try:
            body = response.json()
        except JSONDecodeError:
            content_type = response.headers.get("content-type", "")
            detail = self._response_excerpt(response)
            raise ValueError(f"模型接口返回非 JSON 响应：endpoint={endpoint}，content-type={content_type}，响应={detail}") from None
        if not isinstance(body, dict):
            raise ValueError(f"模型接口返回格式不正确：endpoint={endpoint}")
        return body

    def _response_excerpt(self, response: httpx.Response, limit: int = 300) -> str:
        text = response.text.strip()
        if not text:
            return "<empty>"
        return text[:limit]

    def _chat_completions_endpoint(self, base_url: str) -> str:
        base_url = base_url.rstrip("/")
        path = urlparse(base_url).path.rstrip("/")
        if path.endswith("/chat/completions"):
            return base_url
        if not path:
            return f"{base_url}/v1/chat/completions"
        return f"{base_url}/chat/completions"

    def _merge_generation_params(self, payload: dict, model_params: dict) -> None:
        mapping = {
            "temperature": "temperature",
            "topP": "top_p",
            "presencePenalty": "presence_penalty",
            "frequencyPenalty": "frequency_penalty",
            "maxTokens": "max_tokens",
        }
        for source, target in mapping.items():
            value = model_params.get(source)
            if value is not None:
                payload[target] = value
        extra_params = model_params.get("extraParams")
        if isinstance(extra_params, dict):
            payload.update(extra_params)
